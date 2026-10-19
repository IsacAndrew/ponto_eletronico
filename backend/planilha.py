from calendar import monthrange
from collections import defaultdict
from copy import copy
from datetime import date, datetime, timedelta
from io import BytesIO
from pathlib import Path
from zoneinfo import ZoneInfo

from flask import Blueprint, request, send_file
from openpyxl import load_workbook

from .banco_dados import criar_conexao
from .estrutura_registros import garantir_indices_registros
from .jornada import garantir_estrutura_jornada
from .permissoes import exigir_tipos_conta, renderizar_pagina

rota_planilha = Blueprint("planilha", __name__)
FUSO_HORARIO = ZoneInfo("America/Sao_Paulo")
TIPOS_PERIODO = {"dia", "semana", "mes", "ano"}


def buscar_usuarios_exportacao():
    with criar_conexao() as conexao:
        return conexao.execute(
            """
            SELECT id, nome_completo, tipo_conta, ativo
            FROM usuarios
            ORDER BY nome_completo
            """
        ).fetchall()


def converter_data(valor):
    if not valor:
        return None
    for formato in ("%d/%m/%y", "%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(valor, formato).date()
        except (TypeError, ValueError):
            pass
    return None


def obter_periodo(formulario):
    tipo = formulario.get("periodo", "dia")
    if tipo not in TIPOS_PERIODO:
        return None, None, "Selecione um tipo de período válido."

    if tipo == "dia":
        inicio = converter_data(formulario.get("data_inicio"))
        fim = converter_data(formulario.get("data_fim"))
        if not inicio or not fim:
            return None, None, "Informe as datas inicial e final."
        if fim < inicio:
            return None, None, "A data final deve ser igual ou posterior à data inicial."
        return inicio, fim, None

    if tipo == "semana":
        referencia = converter_data(formulario.get("data_referencia"))
        if not referencia:
            return None, None, "Informe uma data para localizar a semana."
        inicio = referencia - timedelta(days=(referencia.weekday() + 1) % 7)
        return inicio, inicio + timedelta(days=6), None

    if tipo == "mes":
        try:
            valor_mes = formulario.get("mes", "")
            if "/" in valor_mes:
                mes, ano = (int(parte) for parte in valor_mes.split("/"))
            else:
                ano, mes = (int(parte) for parte in valor_mes.split("-"))
            inicio = date(ano, mes, 1)
        except (TypeError, ValueError):
            return None, None, "Informe um mês válido."
        return inicio, date(ano, mes, monthrange(ano, mes)[1]), None

    try:
        ano = int(formulario.get("ano", ""))
        return date(ano, 1, 1), date(ano, 12, 31), None
    except (TypeError, ValueError):
        return None, None, "Informe um ano válido."


def buscar_registros(inicio, fim, ids_usuarios=None):
    garantir_indices_registros()
    consulta = """
        SELECT registro.usuario_id,
               COALESCE(usuario.nome_completo, registro.usuario_nome, 'Usuário removido'),
               COALESCE(usuario.tipo_conta, 'Acesso removido'),
               registro.data_ponto,
               registro.tipo_batida,
               registro.registrado_em
        FROM registros_ponto AS registro
        LEFT JOIN usuarios AS usuario ON usuario.id = registro.usuario_id
        WHERE registro.data_ponto BETWEEN %s AND %s
    """
    parametros = [inicio, fim]
    if ids_usuarios is not None:
        if not ids_usuarios:
            return []
        consulta += " AND registro.usuario_id = ANY(%s)"
        parametros.append(ids_usuarios)
    consulta += " ORDER BY registro.data_ponto, registro.usuario_nome, registro.registrado_em"
    with criar_conexao() as conexao:
        with conexao.cursor() as cursor:
            cursor.execute(consulta, parametros)
            return cursor.fetchall()


def buscar_jornadas(ids_usuarios=None):
    garantir_estrutura_jornada()
    consulta = """
        SELECT jornada.usuario_id, jornada.dia_semana,
               jornada.entrada, jornada.intervalo,
               jornada.retorno_intervalo, jornada.saida,
               usuario.quantidade_batidas
        FROM jornadas_usuarios AS jornada
        JOIN usuarios AS usuario ON usuario.id = jornada.usuario_id
        WHERE TRUE
    """
    parametros = []
    if ids_usuarios is not None:
        if not ids_usuarios:
            return {}
        consulta += " AND jornada.usuario_id = ANY(%s)"
        parametros.append(ids_usuarios)
    with criar_conexao() as conexao:
        with conexao.cursor() as cursor:
            cursor.execute(consulta, parametros)
            registros = cursor.fetchall()

    jornadas = {}
    for id_usuario, dia, entrada, intervalo, retorno, saida, quantidade in registros:
        jornadas.setdefault(id_usuario, {})[dia] = (
            entrada, intervalo, retorno, saida, quantidade
        )
    return jornadas


DIAS_SEMANA = ("Seg", "Ter", "Qua", "Qui", "Sex", "Sáb", "Dom")
MESES = (
    "Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho",
    "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro",
)


def iterar_meses(inicio, fim):
    ano, mes = inicio.year, inicio.month
    while (ano, mes) <= (fim.year, fim.month):
        primeiro_dia = date(ano, mes, 1)
        ultimo_dia = date(ano, mes, monthrange(ano, mes)[1])
        yield primeiro_dia, ultimo_dia
        mes += 1
        if mes > 12:
            mes = 1
            ano += 1


def nome_aba_unico(nome, inicio_mes, titulos_existentes):
    base = f"{nome} {inicio_mes.month:02d}-{inicio_mes.year}"
    base = "".join("_" if caractere in "[]:*?/\\" else caractere for caractere in base)
    base = base.strip("'") or "Folha de Frequência"
    titulo = base[:31]
    contador = 2
    while titulo in titulos_existentes:
        sufixo = f" ({contador})"
        titulo = f"{base[:31 - len(sufixo)]}{sufixo}"
        contador += 1
    return titulo


def preencher_folha(planilha, usuario, jornada, presencas, inicio, fim, inicio_mes):
    id_usuario, nome, tipo_conta, ativo = usuario
    planilha["C2"] = nome
    planilha["D3"] = "Ativo" if ativo else "Inativo"
    planilha["J3"] = tipo_conta
    planilha["U5"] = f"{MESES[inicio_mes.month - 1]} / {inicio_mes.year}"

    dias_jornada = sorted(jornada)
    if dias_jornada:
        planilha["R3"] = ", ".join(DIAS_SEMANA[dia] for dia in dias_jornada)

    dias_com_presenca = presencas.get(id_usuario, set())
    hoje = datetime.now(FUSO_HORARIO).date()
    total_dias_mes = monthrange(inicio_mes.year, inicio_mes.month)[1]
    for numero_dia in range(1, 32):
        linha = 11 + numero_dia
        if numero_dia > total_dias_mes:
            planilha.row_dimensions[linha].hidden = True
            continue
        dia = date(inicio_mes.year, inicio_mes.month, numero_dia)
        planilha.cell(linha, 8, DIAS_SEMANA[dia.weekday()])
        if dia < inicio or dia > fim:
            continue
        if dia in dias_com_presenca:
            planilha.cell(linha, 1, nome)
        elif ativo and dia <= hoje and dia.weekday() in jornada:
            planilha.cell(linha, 1, "Ausente")

    for celula in (planilha["C2"], planilha["D3"], planilha["J3"], planilha["R3"], planilha["U5"]):
        alinhamento = copy(celula.alignment)
        alinhamento.shrink_to_fit = True
        celula.alignment = alinhamento

    planilha.sheet_view.showGridLines = False
    planilha.print_area = "A1:Z45"
    planilha.page_setup.fitToWidth = 1
    planilha.page_setup.fitToHeight = 1
    planilha.sheet_properties.pageSetUpPr.fitToPage = True


def criar_planilha(registros, jornadas, usuarios, inicio, fim):
    caminho_modelo = Path(__file__).resolve().parent / "modelos" / "xlsc_base.xlsx"
    workbook = load_workbook(caminho_modelo)
    modelo = workbook.active
    titulos_existentes = set()
    presencas = defaultdict(set)
    for registro in registros:
        if registro[0] is not None:
            presencas[registro[0]].add(registro[3])

    for usuario in usuarios:
        id_usuario = usuario[0]
        jornada_usuario = jornadas.get(id_usuario, {})
        for inicio_mes, fim_mes in iterar_meses(inicio, fim):
            planilha = workbook.copy_worksheet(modelo)
            planilha.title = nome_aba_unico(usuario[1], inicio_mes, titulos_existentes)
            titulos_existentes.add(planilha.title)
            preencher_folha(
                planilha,
                usuario,
                jornada_usuario,
                presencas,
                max(inicio, inicio_mes),
                min(fim, fim_mes),
                inicio_mes,
            )

    workbook.remove(modelo)
    arquivo = BytesIO()
    workbook.save(arquivo)
    arquivo.seek(0)
    return arquivo


def dados_formulario(formulario):
    return {
        "periodo": formulario.get("periodo", "dia"),
        "data_referencia": formulario.get("data_referencia", ""),
        "data_inicio": formulario.get("data_inicio", ""),
        "data_fim": formulario.get("data_fim", ""),
        "mes": formulario.get("mes", ""),
        "ano": formulario.get("ano", ""),
        "escopos": formulario.getlist("escopo"),
    }


def filtrar_colaboradores(colaboradores, escopos):
    if escopos == ["todos"]:
        return colaboradores, "Todos os usuários"
    if not escopos or "todos" in escopos:
        return None, ""

    try:
        ids_selecionados = {int(valor.split(":", 1)[1]) for valor in escopos
                            if valor.startswith("usuario:")}
    except (IndexError, ValueError):
        return None, ""
    if not ids_selecionados or len(ids_selecionados) != len(escopos):
        return None, ""

    selecionados = [usuario for usuario in colaboradores if usuario[0] in ids_selecionados]
    if len(selecionados) != len(ids_selecionados):
        return None, ""
    nomes = ", ".join(usuario[1] for usuario in selecionados)
    return selecionados, nomes


@rota_planilha.route("/planilha", methods=["GET", "POST"])
@exigir_tipos_conta("Suporte", "Administração")
def exibir_planilha():
    campos = dados_formulario(request.form if request.method == "POST" else request.args)
    erro = None
    if request.method == "POST":
        inicio, fim, erro = obter_periodo(request.form)
        escopos = campos["escopos"]
        if not escopos:
            erro = "Selecione ao menos um usuário ou todos os usuários."
        if not erro:
            colaboradores = buscar_usuarios_exportacao()
            selecionados, _ = filtrar_colaboradores(colaboradores, escopos)
            if not selecionados:
                erro = "Nenhum usuário corresponde à seleção."
            else:
                ids_usuarios = [usuario[0] for usuario in selecionados]
                registros = buscar_registros(inicio, fim, ids_usuarios)
                jornadas = buscar_jornadas(ids_usuarios)
                arquivo = criar_planilha(registros, jornadas, selecionados, inicio, fim)
                nome_arquivo = f"ponto_{inicio.isoformat()}_{fim.isoformat()}.xlsx"
                return send_file(
                    arquivo,
                    as_attachment=True,
                    download_name=nome_arquivo,
                    mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )

    return renderizar_pagina(
        "planilha.html",
        "planilha",
        colaboradores=buscar_usuarios_exportacao(),
        campos=campos,
        erro_planilha=erro,
    )
